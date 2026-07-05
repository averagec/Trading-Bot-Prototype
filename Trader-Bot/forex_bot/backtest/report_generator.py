"""
Report Generator
Outputs backtest results to a formatted Excel report.
"""

import sys
import logging
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone

logger = logging.getLogger(__name__)

DARK = "1A1A2E"; MID = "16213E"; ACCENT = "0F3460"
GREEN = "00C851"; RED = "FF4444"; GOLD = "FFB300"; WHITE = "FFFFFF"

def _font(bold=False, color=WHITE, size=10):
    return Font(bold=bold, color=color, size=size, name="Arial")
def _fill(c): return PatternFill("solid", start_color=c, fgColor=c)
def _center(): return Alignment(horizontal="center", vertical="center", wrap_text=True)
def _border():
    s = Side(style="thin", color="333366")
    return Border(left=s, right=s, top=s, bottom=s)


class ReportGenerator:

    def __init__(self, output_path="backtest_report.xlsx"):
        self.output_path = output_path

    def generate(self, backtest_results: dict, optimiser_summary: dict):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        self._build_summary(wb, backtest_results, optimiser_summary)
        self._build_results(wb, backtest_results)
        self._build_per_pair(wb, optimiser_summary)
        wb.save(self.output_path)
        logger.info(f"Report saved: {self.output_path}")

    def _build_summary(self, wb, results, summary):
        ws = wb.create_sheet("Summary")
        ws.sheet_properties.tabColor = "009900"
        for col, w in zip("ABCDE", [25, 20, 20, 20, 20]):
            ws.column_dimensions[col].width = w

        ws.merge_cells("A1:E1")
        ws["A1"] = "BACKTEST REPORT - SUMMARY"
        ws["A1"].font = _font(bold=True, size=14, color=GOLD)
        ws["A1"].fill = _fill(DARK)
        ws["A1"].alignment = _center()

        info = [
            ("Run Date",          results.get("run_date", "")[:10]),
            ("Pairs Tested",      results.get("pairs_tested", 0)),
            ("Combos Tested",     results.get("total_combos", 0)),
            ("Valid Results",     summary.get("overall_stats", {}).get("total_valid_combos", 0)),
            ("Avg Win Rate",      f"{summary.get('overall_stats', {}).get('avg_win_rate', 0)}%"),
            ("Avg Profit Factor", summary.get("overall_stats", {}).get("avg_profit_factor", 0)),
            ("Optimal Threshold", f"{summary.get('optimal_threshold', 65)}%"),
        ]

        for row, (label, value) in enumerate(info, 3):
            ws[f"A{row}"] = label
            ws[f"B{row}"] = value
            ws[f"A{row}"].font = _font(color="CCCCCC")
            ws[f"A{row}"].fill = _fill(MID)
            ws[f"B{row}"].font = _font(bold=True)
            ws[f"B{row}"].fill = _fill(ACCENT)
            ws[f"B{row}"].alignment = _center()
            ws[f"B{row}"].border = _border()

        row = len(info) + 5
        ws[f"A{row}"] = "OPTIMAL WEIGHTS"
        ws[f"A{row}"].font = _font(bold=True, color=GOLD)
        ws[f"A{row}"].fill = _fill(DARK)
        row += 1

        for strategy, weight in sorted(summary.get("optimal_weights", {}).items(), key=lambda x: x[1], reverse=True):
            ws[f"A{row}"] = strategy.replace("_", " ").title()
            ws[f"B{row}"] = f"{weight} pts"
            ws[f"A{row}"].font = _font(color="CCCCCC")
            ws[f"A{row}"].fill = _fill(MID)
            ws[f"B{row}"].font = _font(bold=True)
            ws[f"B{row}"].fill = _fill(ACCENT)
            ws[f"B{row}"].alignment = _center()
            row += 1

    def _build_results(self, wb, results):
        ws = wb.create_sheet("All Results")
        ws.sheet_properties.tabColor = "FFB300"

        cols   = ["Symbol", "Strategies", "Trades", "Win Rate %",
                  "Profit Factor", "Total Pips", "Max DD", "Expectancy", "Score"]
        widths = [10, 50, 8, 12, 14, 12, 10, 12, 10]

        ws.merge_cells(f"A1:{get_column_letter(len(cols))}1")
        ws["A1"] = "ALL RESULTS (sorted by combined score)"
        ws["A1"].font = _font(bold=True, size=12, color=GOLD)
        ws["A1"].fill = _fill(DARK)
        ws["A1"].alignment = _center()

        for col_idx, (h, w) in enumerate(zip(cols, widths), 1):
            cell = ws.cell(row=2, column=col_idx, value=h)
            cell.font = _font(bold=True)
            cell.fill = _fill(ACCENT)
            cell.alignment = _center()
            cell.border = _border()
            ws.column_dimensions[get_column_letter(col_idx)].width = w

        ws.freeze_panes = "A3"

        for i, r in enumerate(results.get("all_results", [])[:200], 3):
            wr = r.get("win_rate", 0)
            pf = r.get("profit_factor", 0)
            row_data = [
                r.get("symbol"), " + ".join(r.get("strategies", [])),
                r.get("total_trades"), wr, pf,
                r.get("total_pips"), r.get("max_drawdown"),
                r.get("expectancy"), r.get("combined_score"),
            ]
            bg = MID if i % 2 == 0 else DARK
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=i, column=col_idx, value=value)
                cell.fill = _fill(bg)
                cell.alignment = _center()
                cell.border = _border()
                if col_idx == 4:
                    cell.font = _font(bold=True, color=GREEN if wr >= 55 else RED)
                elif col_idx == 5:
                    cell.font = _font(bold=True, color=GREEN if pf >= 1.5 else RED)
                else:
                    cell.font = _font(color=WHITE)

    def _build_per_pair(self, wb, summary):
        ws = wb.create_sheet("Best Per Pair")
        ws.sheet_properties.tabColor = "0000CC"

        cols   = ["Symbol", "Best Strategies", "Win Rate %",
                  "Profit Factor", "Total Trades", "Expectancy"]
        widths = [10, 55, 12, 14, 12, 14]

        ws.merge_cells(f"A1:{get_column_letter(len(cols))}1")
        ws["A1"] = "BEST COMBO PER PAIR"
        ws["A1"].font = _font(bold=True, size=12, color=GOLD)
        ws["A1"].fill = _fill(DARK)
        ws["A1"].alignment = _center()

        for col_idx, (h, w) in enumerate(zip(cols, widths), 1):
            cell = ws.cell(row=2, column=col_idx, value=h)
            cell.font = _font(bold=True)
            cell.fill = _fill(ACCENT)
            cell.alignment = _center()
            cell.border = _border()
            ws.column_dimensions[get_column_letter(col_idx)].width = w

        ws.freeze_panes = "A3"

        for i, (sym, data) in enumerate(summary.get("best_per_pair", {}).items(), 3):
            wr = data.get("win_rate", 0)
            pf = data.get("profit_factor", 0)
            row_data = [
                sym, " + ".join(data.get("strategies", [])),
                wr, pf, data.get("total_trades"), data.get("expectancy"),
            ]
            bg = MID if i % 2 == 0 else DARK
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=i, column=col_idx, value=value)
                cell.fill = _fill(bg)
                cell.alignment = _center()
                cell.border = _border()
                if col_idx == 3:
                    cell.font = _font(bold=True, color=GREEN if wr >= 55 else RED)
                elif col_idx == 4:
                    cell.font = _font(bold=True, color=GREEN if pf >= 1.5 else RED)
                else:
                    cell.font = _font(color=WHITE)