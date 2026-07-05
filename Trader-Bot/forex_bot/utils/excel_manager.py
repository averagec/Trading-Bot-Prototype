"""
Excel Manager
Handles all read/write operations to the forex bot dashboard workbook.
Tabs: Config, Signals, Trade Log, Dashboard, News
"""

import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side)
from openpyxl.utils import get_column_letter

from config import EXCEL, PAIRS, RISK, MIN_SCORE_TO_TRADE

logger = logging.getLogger(__name__)

# ── Colour palette ─────────────────────────────────────────────────────────────
DARK_BG    = "1A1A2E"
MID_BG     = "16213E"
ACCENT     = "0F3460"
GREEN      = "00C851"
RED        = "FF4444"
GOLD       = "FFB300"
WHITE      = "FFFFFF"
LIGHT_GREY = "E8E8E8"
HEADER_BG  = "0F3460"


def _font(bold=False, color=WHITE, size=10):
    return Font(bold=bold, color=color, size=size, name="Arial")

def _fill(hex_color):
    return PatternFill("solid", start_color=hex_color, fgColor=hex_color)

def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def _thin_border():
    side = Side(style="thin", color="333366")
    return Border(left=side, right=side, top=side, bottom=side)


class ExcelManager:

    def __init__(self, path: str = EXCEL["path"]):
        self.path   = Path(path)
        self.sheets = EXCEL["sheets"]

    # ── Workbook Setup ─────────────────────────────────────────────────────────

    def create_workbook(self):
        """Create the full dashboard workbook from scratch."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default sheet

        self._build_config_sheet(wb)
        self._build_signals_sheet(wb)
        self._build_trade_log_sheet(wb)
        self._build_dashboard_sheet(wb)
        self._build_news_sheet(wb)

        wb.save(self.path)
        logger.info(f"Workbook created: {self.path}")

    def _set_tab_color(self, ws, hex_color):
        ws.sheet_properties.tabColor = hex_color

    # ── Config Sheet ───────────────────────────────────────────────────────────

    def _build_config_sheet(self, wb):
        ws = wb.create_sheet(self.sheets["config"])
        self._set_tab_color(ws, ACCENT)

        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 35

        # Title
        ws.merge_cells("A1:C1")
        ws["A1"] = "⚙️  BOT CONFIGURATION"
        ws["A1"].font      = _font(bold=True, size=14, color=GOLD)
        ws["A1"].fill      = _fill(DARK_BG)
        ws["A1"].alignment = _center()
        ws.row_dimensions[1].height = 32

        headers = [("SETTING", "VALUE", "DESCRIPTION")]
        settings = [
            ("── TRADING PAIRS ──", "", ""),
            ("Active Pairs", ", ".join(PAIRS), "Comma-separated forex symbols"),
            ("", "", ""),
            ("── RISK MANAGEMENT ──", "", ""),
            ("Risk Per Trade (%)", RISK["risk_per_trade_pct"], "% of account balance per trade"),
            ("Max Open Trades", RISK["max_open_trades"], "Maximum simultaneous positions"),
            ("Max Daily Loss (%)", RISK["max_daily_loss_pct"], "Stop trading if daily loss hits this %"),
            ("Min R:R Ratio", RISK["min_rr_ratio"], "Minimum reward:risk to take a trade"),
            ("", "", ""),
            ("── STRATEGY THRESHOLDS ──", "", ""),
            ("Min Score to Trade (%)", MIN_SCORE_TO_TRADE, "Minimum confluence % required"),
            ("", "", ""),
            ("── TIMEFRAMES ──", "", ""),
            ("Entry Timeframe", "M15", "Main entry signal timeframe"),
            ("Confirmation Timeframe", "H1", "Mid-level confirmation"),
            ("Bias Timeframe", "H4", "Higher timeframe trend bias"),
            ("", "", ""),
            ("── NEWS SETTINGS ──", "", ""),
            ("No-Trade Window (mins)", 15, "Minutes before/after news to block trading"),
            ("Impact Filter", "High", "Block on: High impact events only"),
        ]

        row = 2
        for col_idx, header in enumerate(["SETTING", "VALUE", "DESCRIPTION"], 1):
            cell = ws.cell(row=row, column=col_idx, value=header)
            cell.font      = _font(bold=True, color=WHITE)
            cell.fill      = _fill(HEADER_BG)
            cell.alignment = _center()
            cell.border    = _thin_border()
        ws.row_dimensions[row].height = 20
        row += 1

        for setting, value, desc in settings:
            is_section = setting.startswith("──")

            c1 = ws.cell(row=row, column=1, value=setting)
            c2 = ws.cell(row=row, column=2, value=value)
            c3 = ws.cell(row=row, column=3, value=desc)

            if is_section:
                for c in [c1, c2, c3]:
                    c.font      = _font(bold=True, color=GOLD)
                    c.fill      = _fill(MID_BG)
                    c.alignment = _left()
            elif setting == "":
                for c in [c1, c2, c3]:
                    c.fill = _fill(DARK_BG)
            else:
                c1.font = _font(color=LIGHT_GREY);   c1.fill = _fill(DARK_BG); c1.alignment = _left()
                c2.font = _font(bold=True, color=WHITE); c2.fill = _fill(MID_BG); c2.alignment = _center()
                c3.font = _font(color=LIGHT_GREY);   c3.fill = _fill(DARK_BG); c3.alignment = _left()
                for c in [c1, c2, c3]:
                    c.border = _thin_border()

            row += 1

    # ── Signals Sheet ──────────────────────────────────────────────────────────

    def _build_signals_sheet(self, wb):
        ws = wb.create_sheet(self.sheets["signals"])
        self._set_tab_color(ws, "009900")

        cols = ["Timestamp", "Symbol", "Direction", "Probability %",
                "Sweep", "Quarterly", "Candle", "Trend", "FVG",
                "News Blocked", "LLM Decision", "Confidence %", "LLM Reasoning"]
        widths = [20, 10, 10, 14, 12, 12, 18, 16, 16, 14, 14, 14, 50]

        ws.merge_cells("A1:M1")
        ws["A1"] = "📊  LIVE SIGNALS"
        ws["A1"].font = _font(bold=True, size=13, color=GOLD)
        ws["A1"].fill = _fill(DARK_BG)
        ws["A1"].alignment = _center()
        ws.row_dimensions[1].height = 28

        for col_idx, (header, width) in enumerate(zip(cols, widths), 1):
            cell = ws.cell(row=2, column=col_idx, value=header)
            cell.font      = _font(bold=True, color=WHITE)
            cell.fill      = _fill(HEADER_BG)
            cell.alignment = _center()
            cell.border    = _thin_border()
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.row_dimensions[2].height = 20
        ws.freeze_panes = "A3"

    # ── Trade Log Sheet ────────────────────────────────────────────────────────

    def _build_trade_log_sheet(self, wb):
        ws = wb.create_sheet(self.sheets["trade_log"])
        self._set_tab_color(ws, GOLD)

        cols = ["#", "Open Time", "Close Time", "Symbol", "Direction",
                "Volume", "Entry", "SL", "TP", "Close Price",
                "Pips", "P&L ($)", "Status", "Strategy Tags"]
        widths = [5, 20, 20, 10, 10, 8, 10, 10, 10, 12, 8, 10, 10, 40]

        ws.merge_cells("A1:N1")
        ws["A1"] = "📒  TRADE LOG"
        ws["A1"].font      = _font(bold=True, size=13, color=GOLD)
        ws["A1"].fill      = _fill(DARK_BG)
        ws["A1"].alignment = _center()
        ws.row_dimensions[1].height = 28

        for col_idx, (header, width) in enumerate(zip(cols, widths), 1):
            cell = ws.cell(row=2, column=col_idx, value=header)
            cell.font      = _font(bold=True, color=WHITE)
            cell.fill      = _fill(HEADER_BG)
            cell.alignment = _center()
            cell.border    = _thin_border()
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.row_dimensions[2].height = 20
        ws.freeze_panes = "A3"

    # ── Dashboard Sheet ────────────────────────────────────────────────────────

    def _build_dashboard_sheet(self, wb):
        ws = wb.create_sheet(self.sheets["dashboard"])
        self._set_tab_color(ws, RED)

        for col, width in zip("ABCDEF", [18, 20, 18, 20, 18, 20]):
            ws.column_dimensions[col].width = width

        ws.merge_cells("A1:F1")
        ws["A1"] = "🤖  FOREX BOT DASHBOARD"
        ws["A1"].font      = _font(bold=True, size=16, color=GOLD)
        ws["A1"].fill      = _fill(DARK_BG)
        ws["A1"].alignment = _center()
        ws.row_dimensions[1].height = 36

        # Account summary labels
        labels = [
            ("Balance ($)", "B3"), ("Equity ($)", "D3"), ("Open P&L ($)", "F3"),
            ("Open Trades", "B5"), ("Daily P&L ($)", "D5"), ("Bot Status", "F5"),
        ]
        for label, cell_ref in labels:
            col = cell_ref[0]
            row = int(cell_ref[1:])
            ws[f"{chr(ord(col)-1)}{row}"] = label
            ws[f"{chr(ord(col)-1)}{row}"].font      = _font(color=LIGHT_GREY)
            ws[f"{chr(ord(col)-1)}{row}"].fill      = _fill(MID_BG)
            ws[f"{chr(ord(col)-1)}{row}"].alignment = Alignment(horizontal="right", vertical="center")

            ws[cell_ref] = "—"
            ws[cell_ref].font      = _font(bold=True, size=12, color=WHITE)
            ws[cell_ref].fill      = _fill(ACCENT)
            ws[cell_ref].alignment = _center()
            ws[cell_ref].border    = _thin_border()

        # Stats section
        ws.merge_cells("A8:F8")
        ws["A8"] = "PERFORMANCE STATS"
        ws["A8"].font      = _font(bold=True, color=GOLD)
        ws["A8"].fill      = _fill(DARK_BG)
        ws["A8"].alignment = _center()

        stat_labels = [
            ("Total Trades", "B9"), ("Win Rate (%)", "D9"), ("Profit Factor", "F9"),
            ("Best Trade ($)", "B10"), ("Worst Trade ($)", "D10"), ("Avg R:R", "F10"),
        ]
        for label, cell_ref in stat_labels:
            col = cell_ref[0]
            row = int(cell_ref[1:])
            ws[f"{chr(ord(col)-1)}{row}"] = label
            ws[f"{chr(ord(col)-1)}{row}"].font = _font(color=LIGHT_GREY)
            ws[f"{chr(ord(col)-1)}{row}"].fill = _fill(MID_BG)
            ws[f"{chr(ord(col)-1)}{row}"].alignment = Alignment(horizontal="right", vertical="center")
            ws[cell_ref] = '=COUNTA(\'Trade Log\'!A3:A1000)' if label == "Total Trades" else "—"
            ws[cell_ref].font      = _font(bold=True, color=WHITE)
            ws[cell_ref].fill      = _fill(ACCENT)
            ws[cell_ref].alignment = _center()
            ws[cell_ref].border    = _thin_border()

    def _right_align(self):
        return Alignment(horizontal="right", vertical="center")

    # ── News Sheet ─────────────────────────────────────────────────────────────

    def _build_news_sheet(self, wb):
        ws = wb.create_sheet(self.sheets["news"])
        self._set_tab_color(ws, "CC0000")

        cols = ["Time (UTC)", "Country", "Impact", "Event", "Forecast", "Previous", "Actual", "Bot Action"]
        widths = [18, 10, 10, 40, 12, 12, 12, 20]

        ws.merge_cells("A1:H1")
        ws["A1"] = "📰  NEWS & ECONOMIC CALENDAR"
        ws["A1"].font      = _font(bold=True, size=13, color=GOLD)
        ws["A1"].fill      = _fill(DARK_BG)
        ws["A1"].alignment = _center()
        ws.row_dimensions[1].height = 28

        for col_idx, (header, width) in enumerate(zip(cols, widths), 1):
            cell = ws.cell(row=2, column=col_idx, value=header)
            cell.font      = _font(bold=True, color=WHITE)
            cell.fill      = _fill(HEADER_BG)
            cell.alignment = _center()
            cell.border    = _thin_border()
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.freeze_panes = "A3"

    # ── Write Methods ──────────────────────────────────────────────────────────

    def write_signal(self, signal_data: dict):
        """Append a signal row to the Signals sheet."""
        wb = openpyxl.load_workbook(self.path)
        ws = wb[self.sheets["signals"]]

        next_row = ws.max_row + 1
        bd       = signal_data.get("breakdown", {})
        direction = signal_data.get("direction", "NONE")

        row_data = [
            datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M"),
            signal_data.get("symbol", ""),
            direction,
            signal_data.get("probability", 0),
            bd.get("liquidity_sweep", {}).get("signal", "—"),
            bd.get("quarterly_theory", {}).get("signal", "—"),
            bd.get("candle_pattern", {}).get("pattern", "—"),
            bd.get("trend", {}).get("trend", "—"),
            bd.get("fvg", {}).get("fvg_type", "—"),
            "YES" if signal_data.get("news_blocked") else "NO",
            signal_data.get("llm_decision", "—"),
            signal_data.get("llm_confidence", 0),
            signal_data.get("llm_reasoning", ""),
        ]

        # Direction colour coding
        dir_color = GREEN if direction == "BUY" else (RED if direction == "SELL" else GOLD)

        for col_idx, value in enumerate(row_data, 1):
            cell           = ws.cell(row=next_row, column=col_idx, value=value)
            cell.fill      = _fill(MID_BG if next_row % 2 == 0 else DARK_BG)
            cell.border    = _thin_border()
            cell.alignment = _center() if col_idx != 13 else _left()
            cell.font      = _font(color=WHITE)

            if col_idx == 3:  # Direction cell
                cell.font = _font(bold=True, color=dir_color)

        wb.save(self.path)

    def write_trade(self, trade_data: dict):
        """Append a trade to the Trade Log sheet."""
        wb       = openpyxl.load_workbook(self.path)
        ws       = wb[self.sheets["trade_log"]]
        next_row = ws.max_row + 1
        trade_num = next_row - 2  # Subtract header rows

        pnl    = trade_data.get("pnl", 0)
        status = trade_data.get("status", "OPEN")

        row_data = [
            trade_num,
            trade_data.get("open_time", ""),
            trade_data.get("close_time", ""),
            trade_data.get("symbol", ""),
            trade_data.get("direction", ""),
            trade_data.get("volume", 0),
            trade_data.get("entry", 0),
            trade_data.get("sl", 0),
            trade_data.get("tp", 0),
            trade_data.get("close_price", ""),
            trade_data.get("pips", ""),
            pnl,
            status,
            trade_data.get("tags", ""),
        ]

        pnl_color = GREEN if pnl > 0 else (RED if pnl < 0 else WHITE)

        for col_idx, value in enumerate(row_data, 1):
            cell           = ws.cell(row=next_row, column=col_idx, value=value)
            cell.fill      = _fill(MID_BG if next_row % 2 == 0 else DARK_BG)
            cell.border    = _thin_border()
            cell.alignment = _center()
            cell.font      = _font(color=WHITE)

            if col_idx == 12:  # P&L cell
                cell.font = _font(bold=True, color=pnl_color)

        wb.save(self.path)

    def update_dashboard(self, account_info: dict, bot_status: str = "RUNNING"):
        """Update the live account stats on the Dashboard sheet."""
        wb = openpyxl.load_workbook(self.path)
        ws = wb[self.sheets["dashboard"]]

        updates = {
            "B3": account_info.get("balance", 0),
            "D3": account_info.get("equity", 0),
            "F3": account_info.get("profit", 0),
            "F5": bot_status,
        }
        for cell_ref, value in updates.items():
            ws[cell_ref] = value
            ws[cell_ref].font      = _font(bold=True, size=12,
                                           color=GREEN if isinstance(value, (int, float)) and value > 0
                                           else RED if isinstance(value, (int, float)) and value < 0
                                           else WHITE)
            ws[cell_ref].fill      = _fill(ACCENT)
            ws[cell_ref].alignment = _center()

        wb.save(self.path)

    def write_news(self, events: list):
        """Write upcoming news events to the News sheet."""
        wb = openpyxl.load_workbook(self.path)
        ws = wb[self.sheets["news"]]

        # Clear old data (keep header rows)
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
            for cell in row:
                cell.value = None

        impact_colors = {"High": RED, "Medium": GOLD, "Low": "006600"}

        for i, event in enumerate(events, 3):
            impact = event.get("impact", "Low")
            color  = impact_colors.get(impact, WHITE)

            row_data = [
                event.get("time_utc", datetime.now()).strftime("%Y-%m-%d %H:%M") if hasattr(event.get("time_utc"), "strftime") else str(event.get("time_utc", "")),
                event.get("country", ""),
                impact,
                event.get("title", ""),
                event.get("forecast", "—"),
                event.get("previous", "—"),
                event.get("actual", "TBA"),
                "BLOCKED" if impact == "High" else "ALLOWED",
            ]
            for col_idx, value in enumerate(row_data, 1):
                cell           = ws.cell(row=i, column=col_idx, value=value)
                cell.fill      = _fill(MID_BG if i % 2 == 0 else DARK_BG)
                cell.border    = _thin_border()
                cell.alignment = _center() if col_idx != 4 else _left()
                cell.font      = _font(color=color if col_idx == 3 else WHITE,
                                       bold=(col_idx == 3))

        wb.save(self.path)

    def read_config(self) -> dict:
        """Read config values from the Config sheet."""
        if not self.path.exists():
            return {}
        wb     = openpyxl.load_workbook(self.path, data_only=True)
        ws     = wb[self.sheets["config"]]
        config = {}
        for row in ws.iter_rows(min_row=3, values_only=True):
            if row[0] and row[1] and not str(row[0]).startswith("──"):
                config[row[0]] = row[1]
        return config


# Singleton
excel_manager = ExcelManager()
